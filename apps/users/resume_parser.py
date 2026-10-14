import os
import re
import pdfplumber
import docx
from django.core.exceptions import ValidationError

def extract_text_from_pdf(file_obj):
    extracted_pages = []
    try:
        with pdfplumber.open(file_obj) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text(x_tolerance=2, y_tolerance=3)
                if page_text:
                    extracted_pages.append(page_text)
        return "\n".join(extracted_pages)
    except Exception as e:
        raise ValidationError(f"Failed to parse PDF document: {str(e)}")

def extract_text_from_docx(file_obj):
    extracted_blocks = []
    try:
        doc = docx.Document(file_obj)

        for para in doc.paragraphs:
            if para.text.strip():
                extracted_blocks.append(para.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    extracted_blocks.append(" | ".join(row_text))

        return "\n".join(extracted_blocks)
    except Exception as e:
        raise ValidationError(f"Failed to parse DOCX document: {str(e)}")

def clean_and_normalize_text(raw_text: str) -> str:
    if not raw_text:
        return ""
    
    # 1. Remove null bytes and non-printable control characters
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', raw_text)
    # 2. Replace fancy Unicode bullets/symbols with standard hyphens
    text = re.sub(r'[•●▪➢➣➤►❖]', '- ', text)
    # 3. Fix hyphenated line-breaks (e.g., "develop-\nment" -> "development")
    text = re.sub(r'(\w+)-\n(\w+)', r'\1\2', text)
    # 4. Collapse multiple horizontal spaces/tabs into a single space
    text = re.sub(r'[ \t]+', ' ', text)
    # 5. Normalize multiple blank lines into a maximum of two newlines (paragraph break)
    text = re.sub(r'\n\s*\n+', '\n\n', text)
    # 6. Strip leading/trailing whitespace from each individual line
    lines = [line.strip() for line in text.split('\n')]
    
    return "\n".join(lines).strip()


def process_resume_file(uploaded_file):
    ext = os.path.splitext(uploaded_file.name)[1].lower()
    uploaded_file.seek(0)

    if ext == ".pdf":
        raw_text = extract_text_from_pdf(uploaded_file)
    elif ext == ".docx":
        raw_text = extract_text_from_docx(uploaded_file)
    else:
        raise ValidationError("Unsupported file format. Only .pdf and .docx are allowed.")

    uploaded_file.seek(0)
    cleaned_text = clean_and_normalize_text(raw_text)

    return {
        "filename": uploaded_file.name,
        "extension": ext,
        "character_count": len(cleaned_text),
        "word_count": len(cleaned_text.split()),
        "raw_text_preview": raw_text[:300] + "..." if len(raw_text) > 300 else raw_text,
        "cleaned_text": cleaned_text
    }

